"""Convierte la lista de segmentos en archivos pensados para estudiar.

- .docx: documento de Word con título, datos de la clase y los párrafos (con
         el minuto en gris al inicio de cada uno). Usa Calibri, que en Antü se
         ve con Carlito (métricamente idéntica), así que el documento queda
         igual en Word (Windows o Wine) y en LibreOffice.
- .txt: texto corrido en párrafos, sin marcas de tiempo (para leer/resumir).
- .md:  mismos párrafos con la marca de tiempo al inicio de cada uno, para
        volver al minuto exacto del video si algo no se entiende.
- .srt: subtítulos estándar (sirven para cargarlos en VLC junto al video).
"""

import re
from datetime import datetime
from pathlib import Path
from typing import Iterable, List

from .segmentos import Segmento, formato_tiempo

FORMATOS_VALIDOS = ("docx", "txt", "md", "srt")
FORMATOS_POR_DEFECTO = ["docx", "txt"]

# Un párrafo nuevo arranca si hubo un silencio largo o si el actual ya es largo
# y la oración terminó. Así el texto queda en bloques legibles, no en una pared.
PAUSA_PARRAFO = 2.5  # segundos
LARGO_PARRAFO = 700  # caracteres


def armar_parrafos(segmentos: Iterable[Segmento]) -> List[Segmento]:
    parrafos: List[Segmento] = []
    actual = None
    for seg in segmentos:
        texto = seg.texto.strip()
        if not texto:
            continue
        if actual is None:
            actual = Segmento(seg.inicio, seg.fin, texto)
            continue
        pausa = seg.inicio - actual.fin
        termino_oracion = actual.texto.rstrip().endswith((".", "?", "!", "…"))
        if pausa >= PAUSA_PARRAFO or (len(actual.texto) >= LARGO_PARRAFO and termino_oracion):
            parrafos.append(actual)
            actual = Segmento(seg.inicio, seg.fin, texto)
        else:
            actual.texto += " " + texto
            actual.fin = seg.fin
    if actual is not None:
        parrafos.append(actual)
    return parrafos


def nombre_seguro(titulo: str) -> str:
    """Convierte un título cualquiera en un nombre de archivo válido en Windows/Linux/macOS."""
    limpio = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", titulo).strip().strip(".")
    limpio = re.sub(r"\s+", " ", limpio)
    return limpio[:120] or "transcripcion"


def _encabezado(origen: str, duracion: float) -> List[str]:
    return [
        f"Fuente: {origen}",
        f"Duración: {formato_tiempo(duracion)}",
        f"Transcripto: {datetime.now():%Y-%m-%d %H:%M}",
    ]


def texto_txt(segmentos, titulo, origen, duracion) -> str:
    lineas = [titulo, "=" * len(titulo), *_encabezado(origen, duracion), ""]
    lineas += [p.texto + "\n" for p in armar_parrafos(segmentos)]
    return "\n".join(lineas)


def texto_md(segmentos, titulo, origen, duracion) -> str:
    lineas = [f"# {titulo}", ""]
    lineas += [f"- {l}" for l in _encabezado(origen, duracion)]
    lineas.append("")
    for p in armar_parrafos(segmentos):
        lineas.append(f"**[{formato_tiempo(p.inicio)}]** {p.texto}")
        lineas.append("")
    return "\n".join(lineas)


def texto_srt(segmentos) -> str:
    bloques = []
    numero = 0
    for seg in segmentos:
        texto = seg.texto.strip()
        if not texto:
            continue
        numero += 1
        bloques.append(
            f"{numero}\n{formato_tiempo(seg.inicio, True)} --> {formato_tiempo(seg.fin, True)}\n{texto}\n"
        )
    return "\n".join(bloques)


def escribir_docx(ruta: Path, segmentos, titulo, origen, duracion) -> None:
    try:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.shared import Pt, RGBColor
    except ImportError as e:
        raise RuntimeError("Falta python-docx. Instalalo con: pip install -U python-docx") from e

    doc = Document()
    doc.core_properties.title = titulo
    doc.core_properties.comments = f"Transcripción de: {origen}"

    for nombre in ("Normal", "Title", "Heading 1"):
        estilo = doc.styles[nombre]
        estilo.font.name = "Calibri"
        rpr = estilo.element.get_or_add_rPr()
        # Fuente también para tildes/eñes y otros alfabetos, e idioma español
        # para que el corrector ortográfico de Word/LibreOffice no marque todo.
        rfonts = rpr.get_or_add_rFonts()
        for atributo in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(atributo), "Calibri")
        for atributo in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(atributo), None)  # el tema pisaría la fuente elegida
        lang = rpr.find(qn("w:lang"))
        if lang is None:
            lang = rpr.makeelement(qn("w:lang"), {})
            rpr.append(lang)
        lang.set(qn("w:val"), "es-AR")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    doc.add_heading(titulo, level=0)
    gris = RGBColor(0x76, 0x76, 0x76)
    for linea in _encabezado(origen, duracion):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(linea)
        run.font.size = Pt(9)
        run.font.color.rgb = gris
    doc.add_paragraph()

    for parrafo in armar_parrafos(segmentos):
        p = doc.add_paragraph()
        marca = p.add_run(f"[{formato_tiempo(parrafo.inicio)}]  ")
        marca.font.size = Pt(8)
        marca.font.color.rgb = gris
        p.add_run(parrafo.texto)
    doc.save(str(ruta))


def guardar(segmentos: List[Segmento], titulo: str, origen: str, duracion: float,
            carpeta: Path, formatos: Iterable[str]) -> List[Path]:
    carpeta.mkdir(parents=True, exist_ok=True)
    base = carpeta / nombre_seguro(titulo)
    escritos = []
    for fmt in formatos:
        ruta = base.with_name(base.name + "." + fmt)
        if fmt == "docx":
            escribir_docx(ruta, segmentos, titulo, origen, duracion)
            escritos.append(ruta)
            continue
        if fmt == "txt":
            contenido = texto_txt(segmentos, titulo, origen, duracion)
        elif fmt == "md":
            contenido = texto_md(segmentos, titulo, origen, duracion)
        elif fmt == "srt":
            contenido = texto_srt(segmentos)
        else:
            raise ValueError(f"Formato desconocido: {fmt}")
        ruta.write_text(contenido, encoding="utf-8")
        escritos.append(ruta)
    return escritos
